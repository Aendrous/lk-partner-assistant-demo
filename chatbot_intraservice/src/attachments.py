# -*- coding: utf-8 -*-
"""Вложения заявки: скрины (VL), письма .msg/.eml, документы PDF/DOC/DOCX → смысл через IEK LLM."""
from __future__ import annotations

import base64
import io
import json
import os
import re
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any

import requests
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
MAIL_EXTS = {".msg", ".eml"}
DOC_EXTS = {".pdf", ".doc", ".docx"}
MAX_FILES = 3
MAX_BYTES = 8 * 1024 * 1024
MAX_DOC_TEXT = 12_000
MIN_PDF_TEXT_CHARS = 80
PDF_VL_MAX_PAGES = 2
PDF_TEXT_MAX_PAGES = 8


def _split_file_ids(file_ids: Any) -> list[str]:
    raw = (file_ids or "").strip() if file_ids is not None else ""
    if not raw:
        return []
    return [p.strip() for p in str(raw).split(",") if p.strip()]


def _split_file_names(files: Any) -> list[str]:
    raw = (files or "").strip() if files is not None else ""
    if not raw:
        return []
    return [p.strip() for p in str(raw).split(",") if p.strip()]


def _ext(name: str) -> str:
    return Path(name or "").suffix.lower()


def _kind_for(name: str, data: bytes) -> str:
    ext = _ext(name)
    if ext in IMAGE_EXTS:
        return "image"
    if ext in DOC_EXTS:
        return "document"
    if data[:4] == b"%PDF":
        return "document"
    if ext in MAIL_EXTS:
        return "mail"
    if data[:2] == b"PK":
        # Office Open XML (.docx, .xlsx) — для .doc иногда zip-контейнер
        if ext in {".docx", ".doc", ""}:
            return "document"
    # magic OLE
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        if ext in {".doc"}:
            return "document"
        return "mail"  # OLE без расширения → часто .msg
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n" or data[:4] == b"RIFF":
        return "image"
    if b"From:" in data[:2000] and b"Subject:" in data[:4000]:
        return "mail"
    return "other"


def _utf16_runs_from_blob(data: bytes) -> str:
    chunks: list[str] = []
    for m in re.finditer(rb"(?:[\x20-\x7e]\x00){12,}", data):
        try:
            s = m.group(0).decode("utf-16-le").strip()
        except Exception:
            continue
        if len(s) < 12:
            continue
        if re.search(r"(?i)^(__|\{|\\\\|Microsoft|WordDocument)", s):
            continue
        chunks.append(s)
    seen: set[str] = set()
    out: list[str] = []
    for c in chunks:
        key = c[:80].lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(c)
        if sum(len(x) for x in out) > MAX_DOC_TEXT:
            break
    return "\n".join(out).strip()


def download_task_file(file_id: str) -> bytes:
    import intraservice

    user = intraservice._user()
    password = (os.environ.get("INTRASERVICE_PASSWORD") or "").strip()
    if not password:
        return b""
    resp = requests.get(
        f"{intraservice.api_base()}/taskfile/{file_id}",
        auth=(user, password),
        timeout=120,
        verify=False,
    )
    if resp.status_code >= 400:
        return b""
    return resp.content or b""


def _strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html or "")
    t = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", t)
    t = re.sub(r"(?is)<br\s*/?>", "\n", t)
    t = re.sub(r"(?is)</p>", "\n", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    t = re.sub(r"&nbsp;", " ", t)
    t = re.sub(r"&amp;", "&", t)
    t = re.sub(r"&lt;", "<", t)
    t = re.sub(r"&gt;", ">", t)
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return re.sub(r"[ \t]{2,}", " ", t).strip()


def extract_eml_text(data: bytes) -> str:
    try:
        msg = BytesParser(policy=policy.default).parsebytes(data)
    except Exception:
        return ""
    parts: list[str] = []
    subj = str(msg.get("Subject") or "").strip()
    frm = str(msg.get("From") or "").strip()
    to = str(msg.get("To") or "").strip()
    if subj:
        parts.append(f"Тема: {subj}")
    if frm:
        parts.append(f"От: {frm}")
    if to:
        parts.append(f"Кому: {to}")
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            ctype = (part.get_content_type() or "").lower()
            if ctype == "text/plain":
                try:
                    body = part.get_content()
                except Exception:
                    payload = part.get_payload(decode=True) or b""
                    body = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
                break
            if ctype == "text/html" and not body:
                try:
                    body = _strip_html(str(part.get_content()))
                except Exception:
                    payload = part.get_payload(decode=True) or b""
                    body = _strip_html(
                        payload.decode(part.get_content_charset() or "utf-8", errors="replace")
                    )
    else:
        try:
            body = str(msg.get_content())
            if "html" in (msg.get_content_type() or "").lower():
                body = _strip_html(body)
        except Exception:
            payload = msg.get_payload(decode=True) or b""
            body = payload.decode("utf-8", errors="replace")
    if body:
        parts.append(str(body).strip()[:6000])
    return "\n".join(parts).strip()


def extract_msg_text(data: bytes) -> str:
    """Текст из Outlook .msg: extract_msg если есть, иначе UTF-16/ASCII строки из OLE."""
    try:
        import extract_msg  # type: ignore

        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as tmp:
            tmp.write(data)
            path = tmp.name
        try:
            msg = extract_msg.Message(path)
            parts = []
            if msg.subject:
                parts.append(f"Тема: {msg.subject}")
            if msg.sender:
                parts.append(f"От: {msg.sender}")
            if msg.to:
                parts.append(f"Кому: {msg.to}")
            body = (msg.body or msg.htmlBody or "") or ""
            if isinstance(body, bytes):
                body = body.decode("utf-8", errors="replace")
            if "<html" in body.lower() or "<body" in body.lower():
                body = _strip_html(body)
            if body:
                parts.append(str(body).strip()[:6000])
            msg.close()
            return "\n".join(parts).strip()
        finally:
            try:
                os.unlink(path)
            except OSError:
                pass
    except Exception:
        pass

    return _utf16_runs_from_blob(data)


def extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
    except Exception:
        return ""
    parts: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = (cell.text or "").strip()
                if t:
                    parts.append(t)
    return "\n".join(parts).strip()[:MAX_DOC_TEXT]


def extract_doc_binary_text(data: bytes) -> str:
    """Текст из бинарного .doc (OLE) без MS Word."""
    try:
        import olefile

        ole = olefile.OleFileIO(io.BytesIO(data))
        parts: list[str] = []
        for stream in ("WordDocument", "1Table", "0Table", "Data"):
            if ole.exists(stream):
                try:
                    blob = ole.openstream(stream).read()
                    chunk = _utf16_runs_from_blob(blob)
                    if chunk:
                        parts.append(chunk)
                except Exception:
                    continue
        ole.close()
        text = "\n".join(parts).strip()
        if len(text) >= 40:
            return text[:MAX_DOC_TEXT]
    except Exception:
        pass
    return _utf16_runs_from_blob(data)[:MAX_DOC_TEXT]


def extract_office_document_text(name: str, data: bytes) -> str:
    ext = _ext(name)
    if ext == ".docx" or data[:2] == b"PK":
        text = extract_docx_text(data)
        if text:
            return text
    if ext == ".doc" or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return extract_doc_binary_text(data)
    if data[:2] == b"PK":
        return extract_docx_text(data)
    return ""


def extract_pdf_text(data: bytes, *, max_pages: int = PDF_TEXT_MAX_PAGES) -> str:
    try:
        import pymupdf as fitz

        doc = fitz.open(stream=data, filetype="pdf")
        parts: list[str] = []
        for i in range(min(len(doc), max_pages)):
            page_text = (doc[i].get_text() or "").strip()
            if page_text:
                parts.append(page_text)
        doc.close()
        return "\n".join(parts).strip()[:MAX_DOC_TEXT]
    except Exception:
        return ""


def pdf_pages_as_png(data: bytes, *, max_pages: int = PDF_VL_MAX_PAGES, dpi: int = 150) -> list[bytes]:
    try:
        import pymupdf as fitz

        doc = fitz.open(stream=data, filetype="pdf")
        out: list[bytes] = []
        for i in range(min(len(doc), max_pages)):
            pix = doc[i].get_pixmap(dpi=dpi)
            out.append(pix.tobytes("png"))
        doc.close()
        return out
    except Exception:
        return []


def _parse_json_blob(content: str) -> dict[str, Any] | None:
    raw = (content or "").strip()
    if not raw:
        return None
    if "```" in raw:
        m = re.search(r"```(?:json)?\s*([\s\S]*?)```", raw)
        if m:
            raw = m.group(1).strip()
    try:
        data = json.loads(raw)
        return data if isinstance(data, dict) else None
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", raw)
        if not m:
            return None
        try:
            data = json.loads(m.group(0))
            return data if isinstance(data, dict) else None
        except json.JSONDecodeError:
            return None


def _llm_chat(model: str, messages: list[dict[str, Any]], *, timeout: int = 180) -> str:
    import llm_client

    resp = requests.post(
        f"{llm_client.base_url()}/chat/completions",
        headers=llm_client._headers(),
        json={"model": model, "messages": messages, "temperature": 0.0, "stream": False},
        timeout=timeout,
        verify=False,
    )
    if resp.status_code >= 400:
        return ""
    data = resp.json()
    return (
        (data.get("choices") or [{}])[0].get("message", {}).get("content")
        or (data.get("choices") or [{}])[0].get("text")
        or ""
    ).strip()


def interpret_image(data: bytes, *, filename: str = "") -> dict[str, Any]:
    """VL: смысл скрина + из какого приложения, не сырой OCR-дамп."""
    import llm_client

    ocr_model = os.environ.get("IEK_LLM_OCR_MODEL") or "iek/qwen3.5-122b-vl-int4-nothink"
    b64 = base64.b64encode(data).decode("ascii")
    # mime
    mime = "image/png"
    if data[:3] == b"\xff\xd8\xff":
        mime = "image/jpeg"
    elif filename.lower().endswith(".webp"):
        mime = "image/webp"
    prompt = (
        "Ты помощник HelpDesk IEK. На скриншоте заявки определи:\n"
        "1) из какого приложения/сайта кадр (lk.iek.ru, bp.iek.ru, Outlook, Chrome, 1С, Excel, другое);\n"
        "2) смысл для исполнителя HelpDesk (что не так / что просит клиент) — 1–2 предложения;\n"
        "3) ключевые факты (номера заказов, адреса, email, ошибки UI) — кратко.\n"
        "Ответ СТРОГО JSON:\n"
        '{"app":"...","meaning_ru":"...","key_facts":["..."],"raw_text":"кратко важный текст со скрина"}\n'
        "Не перечисляй весь OCR. Без markdown."
    )
    content = _llm_chat(
        ocr_model,
        [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime};base64,{b64}"},
                    },
                ],
            }
        ],
    )
    parsed = _parse_json_blob(content) or {}
    meaning = str(parsed.get("meaning_ru") or "").strip()
    app = str(parsed.get("app") or "").strip() or "неизвестно"
    facts = [str(x).strip() for x in (parsed.get("key_facts") or []) if str(x).strip()]
    raw = str(parsed.get("raw_text") or "").strip()
    if not meaning and content:
        meaning = content[:400]
    return {
        "ok": bool(meaning),
        "kind": "image",
        "filename": filename,
        "app": app,
        "meaning_ru": meaning[:500],
        "key_facts": facts[:6],
        "raw_text": raw[:1500],
        "model": ocr_model,
    }


def interpret_document_text(text: str, *, filename: str = "", doc_kind: str = "document") -> dict[str, Any]:
    """Смысл PDF/DOC/DOCX через текстовую IEK LLM."""
    import llm_client

    model = (
        os.environ.get("IEK_LLM_SUPPORT_FALLBACK_MODEL")
        or os.environ.get("IEK_LLM_MODEL")
        or "iek/gpt-oss-120b"
    )
    blob = (text or "").strip()[:7000]
    if not blob:
        return {
            "ok": False,
            "kind": "document",
            "filename": filename,
            "app": doc_kind,
            "meaning_ru": "",
            "key_facts": [],
            "raw_text": "",
        }
    prompt = (
        "Ты помощник HelpDesk IEK. Ниже текст вложения — документ (PDF, DOC или DOCX).\n"
        "Сформулируй смысл для исполнителя (1–2 предложения) и ключевые факты "
        "(реквизиты, ИНН, контакты, номера заказов, что просит партнёр).\n"
        "Ответ СТРОГО JSON:\n"
        '{"app":"PDF / Word / документ","meaning_ru":"...","key_facts":["..."],'
        '"raw_text":"краткая выжимка оригинала"}\n'
        "Без markdown.\n\n"
        f"Файл: {filename}\n\n{blob}"
    )
    content = _llm_chat(model, [{"role": "user", "content": prompt}])
    parsed = _parse_json_blob(content) or {}
    meaning = str(parsed.get("meaning_ru") or "").strip() or content[:400]
    app = str(parsed.get("app") or "").strip() or doc_kind
    facts = [str(x).strip() for x in (parsed.get("key_facts") or []) if str(x).strip()]
    raw = str(parsed.get("raw_text") or "").strip() or blob[:800]
    return {
        "ok": bool(meaning),
        "kind": "document",
        "filename": filename,
        "app": app,
        "meaning_ru": meaning[:500],
        "key_facts": facts[:6],
        "raw_text": raw[:1500],
        "model": model,
    }


def interpret_pdf(data: bytes, *, filename: str = "") -> dict[str, Any]:
    """PDF: текстовый слой; если мало текста — VL по страницам (скан)."""
    text = extract_pdf_text(data)
    if len(text.strip()) >= MIN_PDF_TEXT_CHARS:
        item = interpret_document_text(text, filename=filename, doc_kind="PDF")
        item["extracted_chars"] = len(text)
        item["source"] = "pdf_text"
        return item

    pages = pdf_pages_as_png(data)
    if not pages:
        if text.strip():
            item = interpret_document_text(text, filename=filename, doc_kind="PDF")
            item["extracted_chars"] = len(text)
            item["source"] = "pdf_text_short"
            return item
        return {
            "ok": False,
            "kind": "document",
            "filename": filename,
            "app": "PDF",
            "meaning_ru": "",
            "key_facts": [],
            "raw_text": "",
            "error": "не удалось извлечь текст или отрендерить PDF",
        }

    meanings: list[str] = []
    facts: list[str] = []
    raw_parts: list[str] = []
    apps: list[str] = []
    model = ""
    for i, png in enumerate(pages):
        sub = interpret_image(png, filename=f"{filename}#стр{i + 1}")
        model = str(sub.get("model") or model)
        if sub.get("app"):
            apps.append(str(sub["app"]))
        if sub.get("meaning_ru"):
            meanings.append(str(sub["meaning_ru"]))
        for f in sub.get("key_facts") or []:
            facts.append(str(f))
        if sub.get("raw_text"):
            raw_parts.append(str(sub["raw_text"]))
    meaning = " ".join(meanings).strip() or meanings[0] if meanings else ""
    app = apps[0] if apps else "PDF (скан)"
    seen_f: set[str] = set()
    facts_u: list[str] = []
    for f in facts:
        k = f.lower()
        if k not in seen_f:
            seen_f.add(k)
            facts_u.append(f)
    return {
        "ok": bool(meaning),
        "kind": "document",
        "filename": filename,
        "app": app,
        "meaning_ru": meaning[:500],
        "key_facts": facts_u[:6],
        "raw_text": "\n".join(raw_parts).strip()[:1500],
        "model": model,
        "source": "pdf_vl",
        "pages_vl": len(pages),
    }


def interpret_office_document(data: bytes, *, filename: str = "") -> dict[str, Any]:
    ext = _ext(filename)
    doc_kind = "DOCX" if ext == ".docx" else ("DOC" if ext == ".doc" else "Word")
    text = extract_office_document_text(filename, data)
    item = interpret_document_text(text, filename=filename, doc_kind=doc_kind)
    item["extracted_chars"] = len(text)
    item["source"] = "office_text"
    return item


def interpret_mail_text(text: str, *, filename: str = "") -> dict[str, Any]:
    """Смысл письма .msg/.eml через текстовую IEK LLM."""
    import llm_client

    model = (
        os.environ.get("IEK_LLM_SUPPORT_FALLBACK_MODEL")
        or os.environ.get("IEK_LLM_MODEL")
        or "iek/gpt-oss-120b"
    )
    blob = (text or "").strip()[:7000]
    if not blob:
        return {
            "ok": False,
            "kind": "mail",
            "filename": filename,
            "app": "Outlook (.msg/.eml)",
            "meaning_ru": "",
            "key_facts": [],
            "raw_text": "",
        }
    prompt = (
        "Ты помощник HelpDesk IEK. Ниже текст вложения — письмо (.msg/.eml).\n"
        "Сформулируй смысл для исполнителя (1–2 предложения) и ключевые факты.\n"
        "Ответ СТРОГО JSON:\n"
        '{"app":"Outlook / письмо","meaning_ru":"...","key_facts":["..."],'
        '"raw_text":"краткая выжимка оригинала"}\n'
        "Без markdown.\n\n"
        f"Файл: {filename}\n\n{blob}"
    )
    content = _llm_chat(model, [{"role": "user", "content": prompt}])
    parsed = _parse_json_blob(content) or {}
    meaning = str(parsed.get("meaning_ru") or "").strip() or content[:400]
    app = str(parsed.get("app") or "").strip() or "Outlook / письмо"
    facts = [str(x).strip() for x in (parsed.get("key_facts") or []) if str(x).strip()]
    raw = str(parsed.get("raw_text") or "").strip() or blob[:800]
    return {
        "ok": bool(meaning),
        "kind": "mail",
        "filename": filename,
        "app": app,
        "meaning_ru": meaning[:500],
        "key_facts": facts[:6],
        "raw_text": raw[:1500],
        "model": model,
    }


def analyze_task_attachments(
    task_id: str | int,
    *,
    file_ids: list[str] | None = None,
    file_names: list[str] | None = None,
    max_files: int = MAX_FILES,
) -> dict[str, Any]:
    """Скачать вложения заявки и получить смысл через IEK LLM (скрин / письмо / PDF·DOC)."""
    import intraservice

    empty: dict[str, Any] = {
        "ok": True,
        "items": [],
        "text": "",
        "meaning": "",
        "emails": [],
    }
    if file_ids is None or file_names is None:
        task = intraservice.get_task(str(task_id))
        if file_ids is None:
            file_ids = _split_file_ids(task.get("FileIds"))
        if file_names is None:
            file_names = _split_file_names(task.get("Files"))
    ids = list(file_ids or [])
    names = list(file_names or [])
    if not ids:
        return empty

    items: list[dict[str, Any]] = []
    texts: list[str] = []
    meanings: list[str] = []

    for i, fid in enumerate(ids[:max_files]):
        name = names[i] if i < len(names) else f"file_{fid}"
        try:
            data = download_task_file(fid)
        except Exception as exc:
            items.append({"ok": False, "filename": name, "error": str(exc)[:200]})
            continue
        if not data or len(data) > MAX_BYTES:
            items.append({"ok": False, "filename": name, "error": "пустой или слишком большой файл"})
            continue
        kind = _kind_for(name, data)
        if kind == "image":
            item = interpret_image(data, filename=name)
        elif kind == "mail":
            if _ext(name) == ".eml" or (b"From:" in data[:400] and b"Subject:" in data[:800]):
                mail_text = extract_eml_text(data)
            else:
                mail_text = extract_msg_text(data)
                if not mail_text and _ext(name) in {"", ".msg"}:
                    mail_text = extract_eml_text(data)
            item = interpret_mail_text(mail_text, filename=name)
            item["extracted_chars"] = len(mail_text)
        elif kind == "document":
            if _ext(name) == ".pdf" or data[:4] == b"%PDF":
                item = interpret_pdf(data, filename=name)
            else:
                item = interpret_office_document(data, filename=name)
        else:
            items.append({"ok": False, "filename": name, "kind": kind, "error": "тип не поддержан"})
            continue
        items.append(item)
        if item.get("raw_text"):
            texts.append(str(item["raw_text"]))
        if item.get("meaning_ru"):
            if item.get("kind") == "image":
                label = "скрин"
            elif item.get("kind") == "document":
                label = "документ"
            else:
                label = "письмо"
            app = item.get("app") or ""
            meanings.append(
                f"{label} «{name}»"
                + (f" ({app})" if app else "")
                + f": {item['meaning_ru']}"
            )
        for fact in item.get("key_facts") or []:
            texts.append(str(fact))

    combined_text = "\n".join(texts).strip()
    meaning_blob = "\n".join(meanings).strip()
    blob_for_email = f"{combined_text}\n{meaning_blob}"
    emails = [
        e.lower()
        for e in re.findall(
            r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}",
            blob_for_email,
        )
        if not e.lower().endswith(("@iek.ru", "@iek.group"))
    ]
    # unique preserve order
    seen_e: set[str] = set()
    emails_u: list[str] = []
    for e in emails:
        if e not in seen_e:
            seen_e.add(e)
            emails_u.append(e)

    return {
        "ok": True,
        "items": items,
        "text": combined_text[:4000],
        "meaning": meaning_blob[:2000],
        "emails": emails_u,
    }


def format_attachment_lines(att: dict[str, Any] | None) -> list[str]:
    """Строки для скрытого разбора."""
    if not att:
        return []
    items = att.get("items") or []
    lines: list[str] = []
    for it in items:
        if not isinstance(it, dict) or not it.get("ok"):
            continue
        kind = "Скрин" if it.get("kind") == "image" else ("Документ" if it.get("kind") == "document" else "Письмо")
        name = it.get("filename") or ""
        app = (it.get("app") or "").strip()
        meaning = (it.get("meaning_ru") or "").strip()
        if not meaning:
            continue
        head = f"• {kind}"
        if name:
            head += f" «{name}»"
        if app:
            head += f" · приложение: {app}"
        head += f" — {meaning[:280]}"
        lines.append(head)
        for fact in (it.get("key_facts") or [])[:3]:
            f = str(fact).strip()
            if f and f.lower() not in meaning.lower():
                lines.append(f"  ◦ {f[:160]}")
    if not lines and (att.get("meaning") or "").strip():
        lines.append(f"• Вложение — {str(att.get('meaning')).strip()[:320]}")
    return lines
